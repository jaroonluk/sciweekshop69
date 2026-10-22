# -*- coding: utf-8 -*-
"""สร้างคู่มือการใช้งานระบบรับสมัครร้านค้า (ไฟล์ Word)"""
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "_present_apply_screens"
OUT = ROOT / "คู่มือการใช้งานระบบรับสมัครร้านค้า.docx"
OUT_ASCII = ROOT / "user_manual_vendor_apply.docx"
FONT = "TH Sarabun New"
MAX_IMG_WIDTH = 1400


def set_run_font(run, size=16, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, *, size=16):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    # clear default run and set Thai font
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=size)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size)
    return p


def add_heading_th(doc, text, level=1):
    sizes = {1: 22, 2: 18, 3: 16}
    colors = {
        1: RGBColor(0x0D, 0x8A, 0x7F),
        2: RGBColor(0x1C, 0x16, 0x0E),
        3: RGBColor(0x5C, 0x53, 0x46),
    }
    p = add_para(
        doc,
        text,
        size=sizes.get(level, 16),
        bold=True,
        space_before=14 if level == 1 else 10,
        space_after=6,
        color=colors.get(level),
    )
    return p


def prepare_image(path: Path, max_width=MAX_IMG_WIDTH) -> BytesIO:
    img = Image.open(path)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    w, h = img.size
    if w > max_width:
        ratio = max_width / w
        img = img.resize((max_width, int(h * ratio)), Image.Resampling.LANCZOS)
    buf = BytesIO()
    img.save(buf, format="JPEG", quality=82, optimize=True)
    buf.seek(0)
    return buf


def add_screenshot(doc, filename: str, caption: str, width_cm=15.5):
    path = SHOTS / filename
    if not path.exists():
        add_para(doc, f"[ไม่พบภาพ: {filename}]", size=14, color=RGBColor(0xB4, 0x23, 0x18))
        return
    buf = prepare_image(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(buf, width=Cm(width_cm))
    cap = add_para(
        doc,
        caption,
        size=13,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
        color=RGBColor(0x5C, 0x53, 0x46),
    )
    for run in cap.runs:
        run.italic = True


def set_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(16)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def build():
    doc = Document()
    set_normal_style(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # ----- Cover -----
    add_para(doc, "คณะวิทยาศาสตร์ มหาวิทยาลัยขอนแก่น", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(
        doc,
        "คู่มือการใช้งาน",
        size=28,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=24,
        space_after=4,
        color=RGBColor(0x0D, 0x8A, 0x7F),
    )
    add_para(
        doc,
        "ระบบรับสมัครร้านค้า SCiSQUARE",
        size=22,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_para(
        doc,
        "สำหรับผู้ประกอบการ / ผู้สมัครร้านค้า",
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
        color=RGBColor(0x5C, 0x53, 0x46),
    )
    add_para(
        doc,
        "อธิบายทีละขั้นตอน พร้อมภาพหน้าจอจริงของระบบ",
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=20,
        color=RGBColor(0x5C, 0x53, 0x46),
    )

    # ----- Overview -----
    add_heading_th(doc, "1. ระบบนี้ใช้ทำอะไร", 1)
    add_para(
        doc,
        "ระบบรับสมัครร้านค้า SCiSQUARE เป็นเว็บไซต์สำหรับผู้ประกอบการสมัครคัดเลือกเปิดร้านในพื้นที่คณะวิทยาศาสตร์ มข. "
        "ท่านสามารถกรอกใบสมัคร อัปโหลดเอกสาร และตรวจสอบสถานะใบสมัครได้ด้วยตนเองผ่านเบราว์เซอร์ "
        "ทั้งบนคอมพิวเตอร์และมือถือ",
    )
    add_para(doc, "การใช้งานหลักมี 2 อย่าง:", size=16, bold=True, space_before=6)
    add_bullet(doc, "สมัครใหม่ — กรอกข้อมูล เลือกประเภท/โซน อัปโหลดเอกสาร แล้วส่งใบสมัคร")
    add_bullet(doc, "ตรวจสอบสถานะ — ค้นหาด้วยเบอร์โทรที่ใช้สมัคร เพื่อดูผลตรวจเอกสารและการคัดเลือก")

    add_heading_th(doc, "2. ก่อนเริ่มใช้งาน", 1)
    add_bullet(doc, "ใช้เบราว์เซอร์ที่อัปเดตล่าสุด เช่น Chrome, Edge, Safari หรือ Firefox")
    add_bullet(doc, "เตรียมเบอร์โทรศัพท์ที่ติดต่อได้จริง (ใช้เป็นตัวอ้างอิงค้นหาสถานะภายหลัง)")
    add_bullet(doc, "เตรียมไฟล์เอกสารเป็น JPEG, PNG หรือ PDF ตามที่ระบบระบุในแต่ละช่อง")
    add_bullet(doc, "สมัครได้เฉพาะช่วงที่ระบบแสดงสถานะ “เปิดรับสมัคร” เท่านั้น")
    add_para(
        doc,
        "หมายเหตุ: ภาพในคู่มือนี้ถ่ายจากรอบตัวอย่างที่ปิดรับสมัครแล้ว แต่ขั้นตอนการกรอกแบบฟอร์มเหมือนช่วงเปิดรับ",
        size=14,
        space_before=4,
        color=RGBColor(0x5C, 0x53, 0x46),
    )

    # ----- Home -----
    add_heading_th(doc, "3. หน้าหลักของระบบ", 1)
    add_para(
        doc,
        "เมื่อเปิดลิงก์รับสมัคร ระบบจะแสดงโลโก้คณะ ชื่อโครงการ และสถานะการรับสมัครของรอบนั้น "
        "พร้อมทางเลือกตามสถานะปัจจุบัน",
    )
    add_para(doc, "สิ่งที่ควรดูบนหน้าหลัก", size=16, bold=True, space_before=4)
    add_bullet(doc, "แบนเนอร์สถานะ — แจ้งว่าเปิดรับ / ยังไม่เปิด / ปิดรับสมัครแล้ว พร้อมวัน–เวลา")
    add_bullet(doc, "สมัครใหม่ — ปรากฏเมื่อเปิดรับสมัคร กดเพื่อเริ่มกรอกใบสมัคร")
    add_bullet(doc, "ตรวจสอบสถานะใบสมัคร — ใช้ได้ตลอด เพื่อติดตามผลหลังส่งใบสมัคร")
    add_screenshot(doc, "01_home_gate.png", "ภาพที่ 1 · หน้าหลัก แสดงสถานะรับสมัคร และเมนูตรวจสอบสถานะ")

    add_para(doc, "ตัวอย่างบนมือถือ", size=16, bold=True, space_before=4)
    add_para(
        doc,
        "หน้าจอมือถือจัดเรียงแนวตั้ง เมนูและการ์ดใช้งานเหมือนบนคอมพิวเตอร์ เพียงเลื่อนหน้าจอลงเพื่อดูรายละเอียด",
    )
    add_screenshot(doc, "07_mobile_home.png", "ภาพที่ 2 · หน้าหลักบนมือถือ", width_cm=8.5)

    # ----- Status check -----
    add_heading_th(doc, "4. ตรวจสอบสถานะใบสมัคร", 1)
    add_para(doc, "ใช้เมื่อต้องการดูว่าใบสมัครอยู่ขั้นตอนใด หรือผลตรวจเอกสาร/การคัดเลือกเป็นอย่างไร")
    add_para(doc, "วิธีใช้", size=16, bold=True, space_before=4)
    add_bullet(doc, "จากหน้าหลัก กด “ตรวจสอบสถานะใบสมัคร” หรือ “ตรวจสอบสถานะ →”")
    add_bullet(doc, "กรอกเบอร์โทรที่ใช้สมัคร (บังคับ)")
    add_bullet(doc, "กรอกชื่อ–นามสกุลได้ถ้าต้องการช่วยกรองผล (ไม่บังคับ)")
    add_bullet(doc, "ตอบคำถามป้องกันสแปม (เช่น บวกเลขง่าย ๆ) แล้วกด “ค้นหาสถานะ”")
    add_bullet(doc, "หากต้องการเปลี่ยนข้อคำถาม กด “เปลี่ยนข้อ”")
    add_bullet(doc, "กด “กลับหน้าหลัก” เมื่อต้องการออกจากหน้านี้")
    add_screenshot(doc, "02_status_check.png", "ภาพที่ 3 · หน้าตรวจสอบสถานะใบสมัคร")

    # ----- Apply overview -----
    add_heading_th(doc, "5. สมัครใหม่ — ภาพรวม 4 ขั้นตอน", 1)
    add_para(
        doc,
        "เมื่อเปิดรับสมัคร กด “เริ่มสมัคร” แล้วระบบจะพาเข้าสู่แบบฟอร์ม 4 ขั้นตอน "
        "แถบด้านบนจะบอกว่าอยู่ขั้นไหน สีเข้มคือขั้นปัจจุบัน สีเขียวคือขั้นที่ทำเสร็จแล้ว",
    )
    add_bullet(doc, "ขั้นตอน 1 · ข้อมูล — คุณสมบัติ ชื่อ–นามสกุล เบอร์ติดต่อ และคำถามเสริม (ถ้ามี)")
    add_bullet(doc, "ขั้นตอน 2 · ประเภท/โซน — เลือกประเภทร้านค้า (และโซนตามที่ระบบกำหนด)")
    add_bullet(doc, "ขั้นตอน 3 · เอกสาร — อัปโหลดไฟล์และกรอกข้อความตามรายการ")
    add_bullet(doc, "ขั้นตอน 4 · ยืนยัน — ตรวจทานข้อมูล ตอบคำถามป้องกันสแปม แล้วส่งใบสมัคร")
    add_para(
        doc,
        "ปุ่ม “ออกจากแบบฟอร์ม” ใช้กลับหน้าหลัก (ระบบจะถามยืนยันก่อนออก เพื่อกันข้อมูลหายโดยไม่ตั้งใจ)",
        size=14,
        space_before=4,
        color=RGBColor(0x5C, 0x53, 0x46),
    )

    # ----- Step 1 -----
    add_heading_th(doc, "6. ขั้นตอน 1 · ข้อมูลผู้สมัคร", 1)
    add_para(doc, "กรอกข้อมูลพื้นฐานของผู้สมัครให้ครบก่อนกด “ถัดไป”")
    add_para(doc, "ช่องที่ต้องกรอกโดยทั่วไป", size=16, bold=True, space_before=4)
    add_bullet(doc, "คุณสมบัติของท่าน — เลือกจากรายการ เช่น นิติบุคคล / บุคคลธรรมดา / อื่นๆ")
    add_bullet(doc, "คำนำหน้าชื่อ และชื่อ–นามสกุล ผู้สมัคร")
    add_bullet(doc, "เบอร์ติดต่อ — ใช้เบอร์จริง เพราะต้องใช้ค้นหาสถานะภายหลัง")
    add_bullet(doc, "คำถามเสริมตามรอบ (ถ้ามี) เช่น ความจำเป็นใช้ไฟฟ้ากำลังสูง หรือถังน้ำแข็ง")
    add_para(
        doc,
        "ช่องที่มีเครื่องหมาย * สีแดง คือช่องบังคับ หากกรอกไม่ครบ ระบบจะแจ้งเตือนและยังไปขั้นถัดไปไม่ได้",
        size=14,
        space_before=4,
        color=RGBColor(0x5C, 0x53, 0x46),
    )
    add_screenshot(doc, "03_step1_applicant_info_empty.png", "ภาพที่ 4 · ขั้นตอน 1 แบบฟอร์มว่าง")
    add_screenshot(doc, "03b_step1_applicant_info_sample.png", "ภาพที่ 5 · ขั้นตอน 1 ตัวอย่างการกรอกข้อมูล")
    add_screenshot(doc, "08_mobile_step1.png", "ภาพที่ 6 · ขั้นตอน 1 บนมือถือ", width_cm=8.5)

    # ----- Step 2 -----
    add_heading_th(doc, "7. ขั้นตอน 2 · เลือกประเภท / โซน", 1)
    add_para(
        doc,
        "ในขั้นตอนนี้ให้เลือกประเภทร้านค้า (หรือรายการอาหาร) ตามที่ระบบเปิดให้ในรอบนั้น "
        "บางรอบระบบจะจัดโซนให้อัตโนมัติ บางรอบต้องเลือกโซนก่อนแล้วจึงเลือกประเภท",
    )
    add_para(doc, "แนวทางใช้งาน", size=16, bold=True, space_before=4)
    add_bullet(doc, "อ่านคำอธิบายใต้หัวข้อให้ชัดว่าต้องเลือกอะไรบ้าง")
    add_bullet(doc, "เลือกประเภทร้านจากรายการลงมา")
    add_bullet(doc, "หากระบบให้เลือกโซนด้วย ให้เลือกโซนที่ต้องการก่อนกด “ถัดไป”")
    add_bullet(doc, "กด “ย้อนกลับ” ได้หากต้องการแก้ข้อมูลขั้นที่ 1")
    add_screenshot(doc, "04_step2_zone_category.png", "ภาพที่ 7 · ขั้นตอน 2 เลือกประเภทร้านค้า")
    add_screenshot(doc, "09_mobile_step2.png", "ภาพที่ 8 · ขั้นตอน 2 บนมือถือ", width_cm=8.5)

    # ----- Step 3 -----
    add_heading_th(doc, "8. ขั้นตอน 3 · อัปโหลดเอกสาร", 1)
    add_para(
        doc,
        "อัปโหลดเอกสารตามรายการที่ระบบกำหนดในรอบนั้น แต่ละช่องมีคำอธิบายว่าต้องเตรียมอะไร "
        "และรับไฟล์ชนิดใด (โดยทั่วไป JPEG / PNG / PDF)",
    )
    add_para(doc, "ตัวอย่างเอกสารที่อาจพบ (ขึ้นกับรอบและคุณสมบัติ)", size=16, bold=True, space_before=4)
    add_bullet(doc, "สำเนาหนังสือรับรองการจดทะเบียนนิติบุคคล (ไม่เกินระยะเวลาที่ระบุ และลงนาม/ประทับตรา)")
    add_bullet(doc, "สำเนาบัตรประชาชนของผู้มีอำนาจลงนาม (ยังไม่หมดอายุ และเซ็นรับรองสำเนา)")
    add_bullet(doc, "รายละเอียดอาหารและสินค้า — รายการ ราคา รูปถ่าย แหล่งวัตถุดิบ")
    add_bullet(doc, "แผนบริหารจัดการร้าน — ความสะอาด ของเสีย การบำรุงรักษา มาตรฐาน การตกแต่ง เวลาเปิด–ปิด")
    add_bullet(doc, "ประสบการณ์และความเชี่ยวชาญ — กรอกข้อความในช่องที่ระบบเตรียมไว้")
    add_bullet(doc, "หลักฐานเพิ่มเติม (ถ้ามี) — เกียรติบัตร รางวัล หรือภาพผลงาน")
    add_para(doc, "วิธีอัปโหลด", size=16, bold=True, space_before=6)
    add_bullet(doc, "กด “เลือกไฟล์” ในช่องที่ต้องการ")
    add_bullet(doc, "เลือกรูปหรือ PDF จากเครื่อง แล้วรอให้ระบบแสดงตัวอย่างไฟล์")
    add_bullet(doc, "หากชนิดไฟล์ไม่ถูกต้อง ระบบจะแจ้งให้เลือกใหม่")
    add_bullet(doc, "ช่องบังคับ (*) ต้องมีไฟล์หรือข้อความครบก่อนกด “ถัดไป”")
    add_screenshot(doc, "05_step3_upload_docs.png", "ภาพที่ 9 · ขั้นตอน 3 อัปโหลดเอกสาร")
    add_screenshot(doc, "10_mobile_step3.png", "ภาพที่ 10 · ขั้นตอน 3 บนมือถือ", width_cm=8.5)

    # ----- Step 4 -----
    add_heading_th(doc, "9. ขั้นตอน 4 · ตรวจสอบก่อนส่ง", 1)
    add_para(
        doc,
        "ระบบจะสรุปข้อมูลที่กรอกไว้ให้ตรวจทานอีกครั้ง หากพบข้อผิดพลาด กด “ย้อนกลับ” เพื่อแก้ไข "
        "เมื่อข้อมูลถูกต้องแล้ว ให้ตอบคำถามป้องกันสแปม แล้วกด “ส่งใบสมัคร”",
    )
    add_para(doc, "สิ่งที่ต้องทำก่อนส่ง", size=16, bold=True, space_before=4)
    add_bullet(doc, "อ่านสรุปข้อมูลให้ครบทุกหัวข้อ")
    add_bullet(doc, "ตอบคำถามป้องกันสแปม (เช่น 7 + 5 = ?) ในช่อง “คำตอบ”")
    add_bullet(doc, "กด “เปลี่ยนข้อ” ได้หากต้องการคำถามใหม่")
    add_bullet(doc, "กดปุ่มสีเขียวน้ำทะเล “ส่งใบสมัคร”")
    add_para(
        doc,
        "หลังส่งสำเร็จ ระบบจะแสดงข้อความยืนยัน และสามารถกดไปหน้าตรวจสอบสถานะได้ทันที "
        "แนะนำให้จดหรือเก็บเบอร์โทรที่ใช้สมัครไว้เพื่อติดตามผลภายหลัง",
        size=14,
        space_before=4,
        color=RGBColor(0x5C, 0x53, 0x46),
    )
    add_screenshot(doc, "06_step4_review_submit.png", "ภาพที่ 11 · ขั้นตอน 4 ตรวจสอบก่อนส่งและส่งใบสมัคร")

    # ----- Tips -----
    add_heading_th(doc, "10. ข้อควรระวังและคำถามที่พบบ่อย", 1)

    add_heading_th(doc, "กดสมัครไม่ได้ / ไม่มีปุ่มเริ่มสมัคร", 2)
    add_para(doc, "มักเกิดจากรอบนั้นยังไม่เปิดหรือปิดรับแล้ว ให้ดูแบนเนอร์สถานะด้านบนว่าอยู่ในช่วงเปิดรับหรือไม่")

    add_heading_th(doc, "ไปขั้นตอนถัดไปไม่ได้", 2)
    add_para(doc, "ตรวจว่ากรอกช่องที่มี * ครบหรือยัง และเบอร์โทรอยู่ในรูปแบบที่ถูกต้อง")

    add_heading_th(doc, "อัปโหลดไฟล์ไม่ได้", 2)
    add_para(
        doc,
        "ตรวจชนิดไฟล์ (JPEG / PNG / PDF) และขนาดไฟล์ หากไฟล์ใหญ่เกินไปลองย่อรูปหรือแปลงเป็น PDF แล้วลองใหม่",
    )

    add_heading_th(doc, "ค้นหาสถานะไม่เจอ", 2)
    add_para(
        doc,
        "ใช้เบอร์โทรชุดเดียวกับตอนสมัคร ตอบคำถามป้องกันสแปมให้ถูก "
        "และลองใส่ชื่อ–นามสกุลช่วยกรอง หากยังไม่พบ อาจยังไม่ถึงเวลาประมวลผลหรือเบอร์ไม่ตรง",
    )

    add_heading_th(doc, "สมัครบนมือถือได้ไหม", 2)
    add_para(doc, "ได้ ระบบรองรับมือถือ โดยขั้นตอนและปุ่มเหมือนบนคอมพิวเตอร์")

    # ----- Contact footer -----
    add_heading_th(doc, "11. ติดต่อสอบถาม", 1)
    add_para(
        doc,
        "หากพบปัญหาการใช้งานหรือต้องการข้อมูลเพิ่มเติมเกี่ยวกับการรับสมัคร "
        "กรุณาติดต่อเจ้าหน้าที่ตามช่องทางที่คณะวิทยาศาสตร์ มข. / โครงการ SCiSQUARE ประกาศไว้ในแต่ละรอบ",
    )
    add_para(
        doc,
        "— จบคู่มือ —",
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=18,
        color=RGBColor(0x5C, 0x53, 0x46),
    )

    doc.save(OUT)
    doc.save(OUT_ASCII)
    print(f"saved: {OUT}")
    print(f"saved: {OUT_ASCII}")
    print(f"size: {OUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    build()
