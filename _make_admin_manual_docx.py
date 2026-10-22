# -*- coding: utf-8 -*-
"""สร้างคู่มือการใช้งานระบบ SCI Shop ฝั่ง Admin (ไฟล์ Word)"""
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "_present_admin_screens"
OUT = ROOT / "คู่มือการใช้งานระบบAdmin_SCI_Shop.docx"
OUT_ASCII = ROOT / "user_manual_admin_sci_shop.docx"
FONT = "TH Sarabun New"
MAX_IMG_WIDTH = 1400


def set_run_font(run, size=16, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
             space_after=6, space_before=0, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, *, size=16):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=size)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size)
    return p


def add_heading_th(doc, text, level=1):
    sizes = {1: 22, 2: 18, 3: 16}
    colors = {
        1: RGBColor(0x0D, 0x8A, 0x7F),
        2: RGBColor(0x1C, 0x16, 0x0E),
        3: RGBColor(0x5C, 0x53, 0x46),
    }
    return add_para(
        doc, text,
        size=sizes.get(level, 16), bold=True,
        space_before=14 if level == 1 else 10,
        space_after=6,
        color=colors.get(level),
    )


def prepare_image(path: Path, max_width=MAX_IMG_WIDTH) -> BytesIO:
    img = Image.open(path)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    w, h = img.size
    if w > max_width:
        ratio = max_width / w
        img = img.resize((max_width, int(h * ratio)), Image.Resampling.LANCZOS)
    buf = BytesIO()
    img.save(buf, format="JPEG", quality=82, optimize=True)
    buf.seek(0)
    return buf


def add_screenshot(doc, filename: str, caption: str, width_cm=15.2):
    path = SHOTS / filename
    if not path.exists():
        add_para(doc, f"[ไม่พบภาพ: {filename}]", size=14, color=RGBColor(0xB4, 0x23, 0x18))
        return
    buf = prepare_image(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(buf, width=Cm(width_cm))
    cap = add_para(
        doc, caption, size=13,
        align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10,
        color=RGBColor(0x5C, 0x53, 0x46),
    )
    for r in cap.runs:
        r.italic = True


def set_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(16)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def build():
    doc = Document()
    set_normal_style(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # Cover
    add_para(doc, "คณะวิทยาศาสตร์ มหาวิทยาลัยขอนแก่น", size=16,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "คู่มือการใช้งาน", size=28, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=4,
             color=RGBColor(0x0D, 0x8A, 0x7F))
    add_para(doc, "ระบบ SCI Shop · ฝั่งผู้ดูแล / กรรมการ / การเงิน", size=20, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "อธิบายเมนูและงานที่ทำได้ผ่านระบบ พร้อมภาพหน้าจอจริง", size=14,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18,
             color=RGBColor(0x5C, 0x53, 0x46))

    # 1. Overview
    add_heading_th(doc, "1. ระบบนี้ใช้ทำอะไร (ฝั่ง Admin)", 1)
    add_para(doc,
             "SCI Shop เป็นระบบหลังบ้านสำหรับคณะกรรมการฝ่ายจัดหารายได้และเจ้าหน้าที่ "
             "ใช้ตั้งค่ารอบรับสมัคร ตรวจเอกสารผู้สมัคร คัดเลือกเข้าล็อกร้าน "
             "ติดตามชำระเงิน และพิมพ์ประกาศรายชื่อร้านที่ได้รับการคัดเลือก")
    add_para(doc, "ลำดับงานที่แนะนำบนหน้าหลัก", size=16, bold=True, space_before=6)
    add_bullet(doc, "ตั้งค่าปี · รอบ · โซน · ล็อก")
    add_bullet(doc, "พิจารณาผู้สมัคร (ตรวจเอกสาร / คัดเลือก)")
    add_bullet(doc, "สรุปผู้สมัคร")
    add_bullet(doc, "รายงานร้าน / ล็อก (จัดสรรล็อก)")
    add_bullet(doc, "สถานะชำระเงิน")
    add_bullet(doc, "พิมพ์ Word / Excel ประกาศ")
    add_bullet(doc, "สรุปการใช้งาน (กราฟวิเคราะห์)")

    # 2. Login & roles
    add_heading_th(doc, "2. การเข้าสู่ระบบและสิทธิ์การใช้งาน", 1)
    add_para(doc, "เข้าสู่ระบบด้วย Google บัญชีที่ได้รับสิทธิ์จากผู้ดูแลระบบเท่านั้น")
    add_bullet(doc, "เปิดหน้า login ของระบบ แล้วกดเข้าสู่ระบบด้วย Google")
    add_bullet(doc, "หลังล็อกอินสำเร็จจะเข้าหน้าหลัก SCI Shop")
    add_bullet(doc, "มุมขวาบนแสดงชื่อ อีเมล บทบาท และปุ่ม “ออก”")
    add_para(doc, "บทบาทในระบบ", size=16, bold=True, space_before=6)
    add_bullet(doc, "ผู้ดูแลระบบ (admin) — ใช้งานได้ทุกเมนู รวมจัดการผู้ใช้และบันทึกการใช้งาน")
    add_bullet(doc, "กรรมการฝ่ายจัดหารายได้ (committee) — ตั้งค่า ตรวจคัดเลือก รายงาน ประกาศ สรุปการใช้งาน (ไม่มีเมนูผู้ใช้/บันทึก)")
    add_bullet(doc, "เจ้าหน้าที่การเงิน (finance) — เน้นหน้าหลักและเมนูสถานะชำระเงิน")

    # 3. Home
    add_heading_th(doc, "3. หน้าหลักและตัวเลือกปี / รอบ", 1)
    add_para(doc,
             "หน้าหลักแสดงการ์ดเมนูเรียงตามลำดับงาน สถิติรอบปัจจุบัน "
             "(ผู้สมัคร / อนุมัติแล้ว / ล็อกมีร้าน / จำนวนรอบ) และแถบลำดับขั้นตอน")
    add_para(doc, "ก่อนเริ่มงาน ให้เลือกปีและรอบด้านบนเสมอ", size=16, bold=True, space_before=4)
    add_bullet(doc, "ปี — สลับปีกิจกรรม (เช่น ปี 2569)")
    add_bullet(doc, "รอบ — สลับรอบรับสมัครของปีนั้น (แสดงสถานะเปิดรับถ้ามี)")
    add_bullet(doc, "กดชื่อ “SCI Shop” หรือปุ่ม “← หน้าหลัก” เพื่อกลับหน้าหลักจากหน้างานย่อย")
    add_screenshot(doc, "01_home_hub.png", "ภาพที่ 1 · หน้าหลักระบบ SCI Shop (เมนูงานและสถิติรอบ)")

    # 4. Events / rounds
    add_heading_th(doc, "4. ตั้งค่าปี · รอบ · โซน · ล็อก", 1)
    add_para(doc,
             "เมนูนี้ใช้เตรียมโครงสร้างก่อนเปิดรับสมัคร และกำหนดรอบเปิดร้าน/รับสมัคร "
             "เลือกกิจกรรมทางซ้าย แล้วใช้แท็บทางขวาจัดการรายละเอียด")

    add_heading_th(doc, "4.1 ข้อมูลกิจกรรม", 2)
    add_bullet(doc, "สร้าง / แก้ไขชื่อกิจกรรม ปี พ.ศ. รหัส และโปรแกรมรับสมัคร (เช่น SCiSQUARE / SCiWEEK)")
    add_bullet(doc, "ตั้งเป็น “กิจกรรมใช้งาน” ของปีนั้น")
    add_bullet(doc, "กด “+ สร้างปี / กิจกรรม” เมื่อเริ่มปีใหม่")
    add_screenshot(doc, "02_events_overview.png", "ภาพที่ 2 · แท็บข้อมูลกิจกรรม")

    add_heading_th(doc, "4.2 รอบรับสมัคร (กำหนดรอบเปิดร้าน / เปิดรับสมัคร)", 2)
    add_para(doc, "ใช้กำหนดช่วงเวลาที่ผู้ประกอบการสมัครได้ และตัวเลือกของแบบฟอร์มในรอบนั้น")
    add_bullet(doc, "ตั้งวัน–เวลา “เปิดรับ” และ “ปิดรับ”")
    add_bullet(doc, "ติ๊ก “เปิด” เมื่อต้องการให้รอบนั้นรับสมัครจริงบนหน้าเว็บสาธารณะ")
    add_bullet(doc, "ติ๊ก “ถามไฟฟ้า” / “ถามน้ำแข็ง” หากรอบนั้นต้องถามผู้สมัคร")
    add_bullet(doc, "เลือก “วิธีสมัคร” — ให้เลือกโซนก่อน หรือเลือกประเภทร้านแล้วระบบจัดโซนให้อัตโนมัติ")
    add_bullet(doc, "กด “บันทึก” หลังแก้ค่า")
    add_bullet(doc, "กด “คัดลอกลิงก์ประชาสัมพันธ์” เพื่อนำไปประกาศ (ใช้ได้แม้รอบยังไม่เปิด)")
    add_bullet(doc, "กด “+ เพิ่มรอบ” เมื่อมีรอบถัดไปในปีเดียวกัน")
    add_screenshot(doc, "03_events_rounds.png", "ภาพที่ 3 · แท็บรอบรับสมัคร (กำหนดเปิด–ปิดและตัวเลือกแบบฟอร์ม)")

    add_heading_th(doc, "4.3 โซน", 2)
    add_bullet(doc, "เพิ่ม / แก้ไขรหัสโซน (เช่น A, B, K) ชื่อภาษาไทย และลำดับแสดงผล")
    add_bullet(doc, "กดบันทึกทีละแถวหลังแก้ไข")
    add_screenshot(doc, "04_events_zones.png", "ภาพที่ 4 · แท็บโซน")

    add_heading_th(doc, "4.4 ล็อก", 2)
    add_bullet(doc, "กำหนดล็อกร้าน เช่น รหัสล็อก โซน ประเภทสินค้า/อาหาร จำนวนที่รับ และสถานะใช้งาน")
    add_bullet(doc, "ใช้เป็นฐานในการคัดเลือกผู้สมัครเข้าล็อกในภายหลัง")
    add_screenshot(doc, "05_events_slots.png", "ภาพที่ 5 · แท็บล็อก")

    add_heading_th(doc, "4.5 สลับล็อก", 2)
    add_para(doc, "ใช้เมื่อต้องการสลับผู้คัดเลือกระหว่างล็อก โดยไม่ต้องลบใบสมัคร")
    add_screenshot(doc, "06_events_swap.png", "ภาพที่ 6 · แท็บสลับล็อก")

    add_heading_th(doc, "4.6 คัดลอกโครงสร้าง", 2)
    add_para(doc,
             "คัดลอกโซนและล็อก (และเลือกว่าจะคัดลอกรอบรับสมัครด้วยหรือไม่) "
             "ไปเป็นกิจกรรม/ปีใหม่ โดยไม่คัดลอกผู้สมัคร — เหมาะตอนเปิดปีถัดไป")
    add_screenshot(doc, "07_events_copy.png", "ภาพที่ 7 · แท็บคัดลอกโครงสร้าง")

    # 5. Review
    add_heading_th(doc, "5. พิจารณาผู้สมัคร", 1)
    add_para(doc,
             "งานหลักของกรรมการ: ดูรายชื่อตามโซน เปิดดูเอกสาร บันทึกผลตรวจ "
             "และคัดเลือกเข้าล็อกร้าน")
    add_para(doc, "เครื่องมือด้านซ้าย", size=16, bold=True, space_before=4)
    add_bullet(doc, "ค้นหา — ชื่อ / เบอร์ / ประเภท / รหัสล็อก")
    add_bullet(doc, "โซน — กรองดูทีละโซนหรือทั้งระบบ")
    add_bullet(doc, "สถานะพิจารณา — ทั้งหมด / รอพิจารณา / อนุมัติ / ไม่อนุมัติ / เจ้าเดิม / ค้างชำระปีก่อน")
    add_bullet(doc, "ภาพรวมตรวจเอกสาร — สรุปผ่าน/ไม่ผ่านจากไฟล์แนบอัตโนมัติ")
    add_screenshot(doc, "08_review.png", "ภาพที่ 8 · หน้าพิจารณาผู้สมัคร")

    add_heading_th(doc, "5.1 ตรวจสถานะเอกสาร (Admin)", 2)
    add_bullet(doc, "เปิดการ์ดผู้สมัคร แล้วกดตรวจเอกสาร")
    add_bullet(doc, "ดูผลตรวจอัตโนมัติจากไฟล์แนบ (รูปถ่าย บัตร ทะเบียนบ้าน แผน ฯลฯ)")
    add_bullet(doc, "เลือกผลตรวจ Admin เช่น รอตรวจสอบ / ครบถ้วน / ไม่ครบถ้วน / ขาดคุณสมบัติ")
    add_bullet(doc, "ระบุเอกสารที่ขาดและหมายเหตุการพิจารณาได้")
    add_bullet(doc, "กดบันทึกสถานะ — ระบบไม่ลบรายการผู้สมัคร")
    add_screenshot(doc, "08b_review_status_modal.png", "ภาพที่ 9 · หน้าต่างตรวจสถานะเอกสาร")

    add_heading_th(doc, "5.2 คัดเลือกเข้าล็อก", 2)
    add_bullet(doc, "ตั้งผลการคัดเลือก: ได้รับการคัดเลือก / ไม่ได้รับการคัดเลือก / รอพิจารณา")
    add_bullet(doc, "เมื่อคัดเลือก ให้ระบุล็อกร้านที่ได้รับ")
    add_bullet(doc, "หากล็อกมีผู้คัดเลือกเดิม ระบบจะเปลี่ยนผู้นั้นเป็น “รอพิจารณา” (ไม่ลบรายการ)")
    add_bullet(doc, "การลบใบสมัครทำได้เฉพาะผู้ดูแลระบบ")

    # 6. Summary
    add_heading_th(doc, "6. สรุปผู้สมัคร", 1)
    add_para(doc,
             "ตารางภาพรวมสถานะเอกสารและผลการคัดเลือก กรองตามสถานะ เอกสาร "
             "เจ้าเดิม หรือค้างชำระปีก่อนได้ และดาวน์โหลดประกาศจากหน้านี้อีกทางหนึ่ง")
    add_screenshot(doc, "09_summary.png", "ภาพที่ 10 · หน้าสรุปผู้สมัคร")

    # 7. Report
    add_heading_th(doc, "7. รายงานร้าน / ล็อก", 1)
    add_para(doc,
             "ดูภาพรวมล็อกทั้งหมดว่ามีผู้ประกอบการหรือยัง จัดสรรร้านเข้าล็อก "
             "และดูรายการประเภทที่ยังว่าง เหมาะใช้ตอนจัดผังร้าน")
    add_bullet(doc, "ดูล็อกที่เต็ม / ว่าง ในแต่ละโซน")
    add_bullet(doc, "เลือกผู้สมัครเข้าล็อกจากหน้ารายงานได้โดยตรง")
    add_bullet(doc, "ดูสรุปรายการที่ยังไม่มีผู้สมัคร")
    add_screenshot(doc, "10_report.png", "ภาพที่ 11 · หน้ารายงานร้าน / ล็อก")

    # 8. Finance
    add_heading_th(doc, "8. สถานะชำระเงิน", 1)
    add_para(doc,
             "รายการร้านที่ได้รับการคัดเลือกแล้ว ติ๊กสถานะชำระเงินเพื่อใช้พิจารณาปีต่อไป "
             "และซิงก์รายชื่อเจ้าเดิมที่ค้างชำระไปเตือนในปีถัดไป")
    add_bullet(doc, "ติ๊ก “ชำระแล้ว” / “ยังไม่ชำระ”")
    add_bullet(doc, "ดูบันทึกล่าสุดและหมายเหตุได้")
    add_bullet(doc, "กด “ซิงก์ alumni” เพื่ออัปเดตรายการค้างชำระสำหรับปีถัดไป")
    add_screenshot(doc, "11_finance.png", "ภาพที่ 12 · หน้าสถานะชำระเงิน")

    # 9. Print / export
    add_heading_th(doc, "9. พิมพ์รายงาน / ประกาศ (Word และ Excel)", 1)
    add_para(doc,
             "บนหน้าหลักมีส่วน “พิมพ์ Word ประกาศร้านค้า” "
             "ใช้ดาวน์โหลดรายชื่อร้านที่ได้รับการคัดเลือกเพื่อนำไปประกาศ")
    add_para(doc, "ขั้นตอน", size=16, bold=True, space_before=4)
    add_bullet(doc, "เลือกปีที่ประกาศ")
    add_bullet(doc, "เลือกรอบเปิดร้านค้า")
    add_bullet(doc, "กด “ดาวน์โหลด Word ประกาศ” หรือ “ดาวน์โหลด Excel ประกาศ”")
    add_bullet(doc, "ระบบจะสร้างไฟล์เฉพาะร้านที่สถานะ “ได้รับการคัดเลือก” ในรอบนั้น")
    add_para(doc,
             "ปุ่มดาวน์โหลดเดียวกันมีในหน้าสรุปผู้สมัครและหน้ารายงานร้าน / ล็อกด้วย",
             size=14, space_before=4, color=RGBColor(0x5C, 0x53, 0x46))
    add_screenshot(doc, "15_announce_home.png",
                   "ภาพที่ 13 · ส่วนพิมพ์ประกาศบนหน้าหลัก (Word / Excel)")

    # 10. Users
    add_heading_th(doc, "10. ผู้ใช้งาน (เฉพาะผู้ดูแลระบบ)", 1)
    add_para(doc, "ค้นหาบุคลากรจาก eoffice แล้วกำหนดสิทธิ์เข้าใช้งานระบบ")
    add_bullet(doc, "พิมพ์ชื่อ / นามสกุล / อีเมล / รหัส เพื่อค้นหา")
    add_bullet(doc, "กำหนดสิทธิ์: ผู้ดูแลระบบ / กรรมการ / การเงิน")
    add_bullet(doc, "เปิด–ปิดการใช้งาน หรือเปลี่ยนสิทธิ์ภายหลังได้")
    add_bullet(doc, "ไม่สามารถลดสิทธิ์หรือปิดบัญชีของตนเองได้")
    add_screenshot(doc, "12_users.png", "ภาพที่ 14 · หน้าผู้ใช้งาน")

    # 11. Audit
    add_heading_th(doc, "11. บันทึกการใช้งาน (เฉพาะผู้ดูแลระบบ)", 1)
    add_para(doc, "ดูประวัติการเข้าใช้งานและการดำเนินการสำคัญ เช่น ให้สิทธิ์ เปลี่ยนสิทธิ์ หรือบันทึกสถานะ")
    add_screenshot(doc, "13_audit.png", "ภาพที่ 15 · หน้าบันทึกการใช้งาน")

    # 12. Usage
    add_heading_th(doc, "12. สรุปการใช้งาน", 1)
    add_para(doc,
             "ดูกราฟและสถิติช่วยตัดสินใจ เช่น อันดับประเภทที่สมัครมาก–น้อย "
             "ช่วงเวลาที่ผู้สมัครเข้ามา และวันเปิด–ปิดรับสมัครของรอบ")
    add_screenshot(doc, "14_usage.png", "ภาพที่ 16 · หน้าสรุปการใช้งาน")

    # 13. Cheat sheet
    add_heading_th(doc, "13. สรุปสั้น ๆ — อยากทำอะไร กดเมนูไหน", 1)
    add_bullet(doc, "เปิด/ปิดรับสมัคร หรือกำหนดวันรอบ → ปี · รอบ · โซน · ล็อก → แท็บรอบรับสมัคร")
    add_bullet(doc, "เพิ่มโซน / ล็อก / ประเภท → ปี · รอบ · โซน · ล็อก → แท็บโซนหรือล็อก")
    add_bullet(doc, "ตรวจเอกสารและคัดเลือก → พิจารณาผู้สมัคร")
    add_bullet(doc, "ดูภาพรวมทั้งรอบ → สรุปผู้สมัคร")
    add_bullet(doc, "จัดร้านเข้าล็อก / ดูล็อกว่าง → รายงานร้าน / ล็อก")
    add_bullet(doc, "ติ๊กชำระเงิน → สถานะชำระเงิน")
    add_bullet(doc, "พิมพ์ประกาศ Word/Excel → การ์ดพิมพ์ประกาศบนหน้าหลัก")
    add_bullet(doc, "ให้สิทธิ์ผู้ใช้ใหม่ → ผู้ใช้งาน (admin)")
    add_bullet(doc, "ดูประวัติการใช้งาน → บันทึกการใช้งาน (admin)")

    # 14. Tips
    add_heading_th(doc, "14. ข้อควรระวัง", 1)
    add_bullet(doc, "เลือกปีและรอบให้ถูกก่อนทำงานทุกครั้ง")
    add_bullet(doc, "ติ๊ก “เปิด” ในรอบรับสมัครเมื่อพร้อมจริงเท่านั้น")
    add_bullet(doc, "การคัดเลือกเข้าล็อกจะกระทบผู้คัดเลือกเดิมของล็อกนั้น (เปลี่ยนเป็นรอพิจารณา)")
    add_bullet(doc, "ประกาศ Word/Excel จะมีเฉพาะร้านที่อนุมัติแล้ว — ถ้ายังเป็น 0 ร้าน ไฟล์จะว่างหรือมีน้อย")
    add_bullet(doc, "ข้อมูลการเงินใช้เตือนเจ้าเดิมปีถัดไป ควรอัปเดตให้ตรงความเป็นจริง")

    add_para(doc, "— จบคู่มือฝั่ง Admin —", size=14,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18,
             color=RGBColor(0x5C, 0x53, 0x46))

    doc.save(OUT)
    doc.save(OUT_ASCII)
    print(f"saved: {OUT}")
    print(f"saved: {OUT_ASCII}")
    print(f"size: {OUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    build()
